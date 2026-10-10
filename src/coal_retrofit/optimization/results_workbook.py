"""结果工作簿 `ccs_results.xlsx`（2026-10-02 起）：表名与行列结构参照 ChinaCCS.xlsm 的结果表，数字全部取自本次求解。

`runner.solve` 写完九张 CSV 后调用，写在同一目录；工作簿出错只记日志并删掉工作簿（删不掉也只记日志），CSV 与
result.json 照写。输入是那九张表（`results_tables.build_result_tables`，键为表名，2026-10-10 前是 15 张、键为文件名）与
CSV 里没有的几样：各年成本类别的折现权重与残值台账（`solution`）、
源与汇所在的节点、节点经纬度、封存汇的每吨封存成本（`prepared`）、节点的省（`results_regions.node_provinces`）。各表的
口径写在工作簿的「说明」表（`_notes`）；源、汇与输送各表在 `results_workbook_sources`，管道与成本各表在
`results_workbook_network`。
"""
from __future__ import annotations

import math
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook

from ..constants import NETWORK_DETOUR_FACTOR
from ._shared import PreparedInputs
from .results_tracing import FLOW_TOL
from .results_workbook_network import cost_lines, cost_sheets, edge_years, pipe_sheets
from .results_workbook_sources import (
    CAPTURE_STOCKS,
    MILLION,
    Block,
    captured_by_year,
    industry_routes,
    plant_pathways,
    sink_results,
    sinks_frame,
    source_results,
    sources_frame,
    stock_table,
    transmission_matrix,
)
from .scenario import OptimizationAssumptions, OptimizationScenario


def _notes(assumptions: OptimizationAssumptions) -> list[list[Any]]:
    """「说明」表：各表的口径；写进去的参数取本次求解的值。"""
    life = int(assumptions.pipeline_lifetime_years)
    credit = float(assumptions.eor_credit_cny_per_t)
    pipe_om, route_opex = float(assumptions.pipe_fixed_om_fraction), float(assumptions.route_opex_cny_per_t_km)
    om_parts = [f"在役管道的投资 x {pipe_om:.1%}/年，闲置的管也计"] if pipe_om else []
    om_parts += [f"按流量计的每吨公里 {route_opex:g} 元"] if route_opex else []
    transport_om = "；".join(om_parts) or "不计"
    storage, offshore = float(assumptions.storage_cost_cny_per_t), float(assumptions.offshore_storage_multiplier)
    return [
        ["项目", "说明"],
        ["来源", "每次求解后由 coal_retrofit.optimization.results_workbook 写出，表名与行列结构参照 ChinaCCS.xlsm 的结果表，"
                "数字全部来自本次求解，与同目录的九张 CSV 是同一组结果。"],
        ["年份", "列名里的年份是规划年。捕集、注入、流量与运行费是该年一年的量，投资是该年一次性的支出。"],
        ["单位", "捕集、注入、输送、流量 Mt/yr；管长、距离 km；成本 百万元（人民币；ChinaCCS 原表为百万美元，这里不换算）；"
                "每吨成本 元/t。"],
        ["分区", "1 东北、2 华北、3 华东、4 西北、5 西南、6 中南，省到区的对应取自 ChinaCCS.xlsm（西藏不在其中，归西南）。"
                "7 海上：海上封存汇与落在省界之外的管网节点。8 跨地理分区：两端在不同陆上分区的管道；一端在海上的管道记陆上"
                "那一端的分区。ChinaCCS 没有海上区（海上汇归沿海的区），跨地理分区记 7：这里另设海上占 7，跨地理分区顺延为 8。"],
        ["源", "煤电 hub（plant_id）与工业 hub（hub_id），省取输入表的省名。部门：电力 = 煤电（捕集含 BECCS），"
              "钢铁 = 高炉-转炉 + 电炉，水泥，化工 = 合成氨 + 甲醇。炼化、现代煤化工、天然气、液化、烯烃、乙二醇不在模型内，"
              "不出对应的表。"],
        ["汇", "本情景纳入的全部封存 hub，没有注入的也列出（storage_scope 为 dsa_only 时不含 EOR 汇）。Sink_Category 取 "
              "storage_type：DSA 深部咸水层、EOR 驱油。海上汇（storage_hubs.csv 的 offshore）的省记 Offshore，其余取 "
              "storage_hubs.csv 的省名，没填的同下条的其余管网节点。"],
        ["节点的省", "源节点取源的省；本情景纳入的封存节点见上，不纳入的（如 dsa_only 去掉的 EOR 汇）同其余管网节点；"
                  "其余管网节点取 pipeline_nodes.csv 自带的省名，没有的按经纬度落在 "
                  "data/ChinaMap/provinces.shp 的哪个省，落在所有省之外的（海上或近岸）记 Offshore。network.csv 的 "
                  "from_province、to_province 与 sinks.csv 的 province 同此。"],
        ["Source_Results", "每个源一行。CO2_emit_rate 是本年的基线排放（不改造时的排放）；CAPEX_Capture 是本年新建捕集能力的"
                           "投资；OPEX_Capture_p、_c、_i、_a 依次是电力、水泥、钢铁、化工源本年的捕集运行费（煤电 = 捕集岛固定"
                           "运维 + CCS 额外燃料，工业 = 捕集路线的年度费）；其后是分部门的基线排放与捕集量。右侧是分区与分省的"
                           "捕集量。"],
        ["Transmission_Matrix", "行是源所在的省，列是汇所在的省（Offshore 为海上汇），值是本年从该省的源送到该省的汇的 CO2，"
                                "零格留空。几个源的 CO2 在管网里汇流以后，哪个源送到哪个汇没有唯一答案；这里按节点充分混合"
                                "追踪：每个节点流出的 CO2（流向下游各边的与在本节点封存的）按流入它的各来源的比例分配（先把"
                                "两个方向的流量相抵，再抵消环流）。行和 = 该省的捕集量，列和 = 该省（或海上）的封存量。"],
        ["*_stock", "CO2_capture_<部门>_stock 是每个源各年的捕集量，CO2_inject_stock 是每个汇各年的注入量；只列有一年不为零的"
                    "行，零格留空。"],
        ["管道", f"管径档是模型的输量档（Mtpa），不折算英寸管径。在役 = 建成不满管道寿命（{life} 年）的整根管，与求解器同口径；"
                "既有走廊容量不是管，不在其中。每条边以 edge_id 区分（同一对节点之间可能有不止一条候选边）；Node1、Node2 是"
                "边的起点、终点，流量是净流量，起点到终点为正。Distance 是两端点的大圆距离，Length_Weighted 是模型用的管长"
                f"（直线距离 x 绕行系数 {NETWORK_DETOUR_FACTOR:g}，或既有走廊的实际路由长；连通性修复补的边（edge_stitch_、"
                "edge_merge_）直线距离不足 1 km 的按 1 km 计，运行期支线（edge_runtime_）管长不足 1 km 的按 1 km 计，这几条边的 "
                "Length_Weighted 可达 Distance 的数倍）。<年>长度 = 在役根数 x 管长。"],
        ["taransport_analysis", "表名照 ChinaCCS 的原拼写。上半部分是各管径档的在役根数与按根累计的管长，下半部分按分区"
                                "汇总。Total_CO2_Pipeline_Length 给两种总管长：有在役管的边每条只算一次，与按根累计。"],
        ["成本", "Cost_* 与 Revenue_EOR 是各年不折现的值，投资计在建成那一年。捕集 = 煤电捕集岛投资 + 工业捕集投资 + 煤电"
                "捕集岛固定运维 + 煤电 CCS 额外燃料（CCS 与 BECCS 多烧的燃料）+ 工业捕集路线的年度费，不含氢路线的费用，也不含"
                f"空冷背压与生物质掺烧的效率损失；运输 = 管道投资 + 运输运维（{transport_om}）；封存 = 扣 EOR 抵扣之前的"
                f"封存费（每吨 {storage:g} 元，海上汇再乘 "
                f"{offshore:g}）；Revenue_EOR = 各汇封存量 x EOR 容量占比 x 每吨 {credit:g} 元的抵扣（参数 eor_credit_cny_per_t，"
                "⚠ 无出处）。"],
        ["Cost_Analysis", "当期成本 = 本年投资 + 本年运行费。每吨运行费 = 本年运行费 ÷ 本年捕集量。全期每吨成本 = 各环节成本的"
                          "现值（投资减去期末残值抵扣，折现与目标函数相同）÷ 捕集量的现值（与年度成本同一折现与年金权重）。"
                          "平均运程 = Σ(各边净流量的绝对值 x 管长) ÷ 捕集量。本年捕集量不超过 "
                          f"{FLOW_TOL:g} Mt（数值噪声）时，该年的每吨运行费与平均运程留空。EOR 抵扣冲减成本，记负。"],
        ["objective", "目标函数值（现值），及各成本类别各年的现值（system.csv 的 cost_discounted_cny），各年各类相加即目标函数值。"
                      "期末残值抵扣（salvage_credit）记在最后一个规划年。"],
        ["本模型另加的表", "Plant_Pathways：煤电 hub 各年各路径的份额；already_air 是原本就是空冷的份额，air_retrofit_operating、"
                         "air_retrofit_installed 是湿冷改空冷的装机里当年在运行的与在役的（建成未满改造寿命），都按占全厂的份额（sources.csv "
                         "煤电行的 air_operating_share、air_installed_share 乘 1 − already_air_share）。Industry_Routes：工业 hub 各年"
                         "各路线的份额（混合原料的甲醇 hub 是份额为正的点源那部分产量的份额，占全厂产量的比例见 sources.csv 的 "
                         "abatable_production_share）。"],
        ["不出的表", "num_total_well_new：模型没有注入井（ChinaCCS 用注入量除以单井注入率算，模型没有单井注入率）。"
                   "Cost_Analysis 的静态投资回收期：CCS 链条除 EOR 抵扣外没有收入（售电收入计在煤电的基线净成本里），算不出"
                   "回收期。ChinaCCS 的「产品收益」（= Revenue_EOR）在本表是「EOR 抵扣」一行，记负。"],
    ]


def write_ccs_workbook(
    path: Path,
    tables: Mapping[str, pd.DataFrame],
    solution: Mapping[str, Any],
    prepared: PreparedInputs,
    scenario: OptimizationScenario,
    assumptions: OptimizationAssumptions,
    node_province: Mapping[str, str],
) -> None:
    """按九张结果表（*tables*，键为表名，`results_tables.TABLES`）与解写 ChinaCCS 版式的工作簿到 *path*。"""
    years = [int(y) for y in scenario.planning_years]
    end_year = years[-1] + scenario.interval_years(tuple(years), len(years) - 1, assumptions)
    tiers = tuple(float(t) for t in assumptions.pipe_capacity_tiers_mtpa)
    sources = sources_frame(tables, prepared)
    sinks = sinks_frame(tables, prepared, node_province)
    edges = edge_years(tables, tiers)
    network = tables["network"]
    flows = network[network["flow_from_node_id"].fillna("").astype(str) != ""]  # 有流量的边

    sheets: dict[str, list[Block]] = {"说明": [(1, 1, _notes(assumptions))]}
    for y in years:
        sheets[f"Source_Results_{y}"] = source_results(sources[sources["year"] == y])
    sheets["Sink_Results"] = sink_results(sinks, years)
    for y in years:
        sheets[f"Transmission_Matrix_{y}"] = [(1, 1, transmission_matrix(
            flows[flows["year"] == y], sources[sources["year"] == y], sinks[sinks["year"] == y], node_province))]
    for key, groups in CAPTURE_STOCKS.items():
        sheets[f"CO2_capture_{key}_stock"] = [(1, 1, stock_table(sources[sources["group"].isin(groups)], "captured_mt",
                                                                 years))]
    sheets["CO2_inject_stock"] = [(1, 1, stock_table(sinks, "use_mt", years))]
    sheets.update(pipe_sheets(edges, years, tiers, prepared, node_province))
    captured = captured_by_year(sources, years)
    sheets.update(cost_sheets(cost_lines(tables, sinks, years), captured, edges, solution, years, end_year))
    sheets["objective"] = _objective(tables["system"], float(solution["objective_cny"]), years)
    sheets["Plant_Pathways"] = [(1, 1, plant_pathways(tables, years))]
    sheets["Industry_Routes"] = [(1, 1, industry_routes(tables, years))]
    _save(Path(path), sheets)


def _objective(breakdown: pd.DataFrame, objective_cny: float, years: Sequence[int]) -> list[Block]:
    """objective：目标函数值，及各成本类别各年的现值（百万元，`system.csv` 的 `cost_discounted_cny`）与合计。"""
    order = list(dict.fromkeys(breakdown["category"]))
    discounted = breakdown.pivot(index="category", columns="year", values="cost_discounted_cny").reindex(
        index=order, columns=list(years)).fillna(0.0) / MILLION
    kinds = breakdown.drop_duplicates("category").set_index("category")["kind"].to_dict()
    rows: list[list[Any]] = [["category", "kind", *years, "合计"]]
    rows += [[category, kinds[category], *row, row.sum()] for category, row in discounted.iterrows()]
    totals = discounted.sum(axis=0)
    rows.append(["合计", None, *totals, totals.sum()])
    return [(1, 1, [["目标函数（现值，百万元）", objective_cny / MILLION]]),
            (3, 1, [["各成本类别各年的现值（百万元）"]]), (4, 1, rows)]


def _save(path: Path, sheets: Mapping[str, list[Block]]) -> None:
    """逐表逐块写格；空值、NaN 不写。"""
    book = Workbook()
    book.remove(book.active)
    for name, blocks in sheets.items():
        sheet = book.create_sheet(name)
        for top, left, rows in blocks:
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    cell = _cell(value)
                    if cell is not None:
                        sheet.cell(row=top + i, column=left + j, value=cell)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _cell(value: Any) -> Any:
    """numpy 标量换成 Python 数，NaN 记空。"""
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    return value
