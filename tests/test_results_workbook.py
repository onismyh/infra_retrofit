"""结果工作簿 `ccs_results.xlsx` 在 toy 上的对账（需要 Gurobi）：输送矩阵的行和、列和；各部门捕集之和等于总捕集；
在役根数乘档容量等于在役的管道容量；不折现成本乘回折现系数等于 `cost_breakdown.csv`；残值台账复算出 `salvage_credit`；
全期每吨成本按定义复算；目标函数各项之和等于目标函数值。工作簿出错时 CSV 与 result.json 照写。"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import pytest

from coal_retrofit import cli, runner
from coal_retrofit.optimization.results_workbook_network import _tier_counts, ledger_salvage
from coal_retrofit.optimization.salvage import remaining_fraction
from test_runner_cli import TABLES, _registry

pytest.importorskip("gurobipy", reason="gurobipy is required for solver integration tests")

ARGV = ["run", "TOY", "--threads", "1", "--time-limit", "60"]
YEARS = (2050, 2060)  # test_runner_cli 的 TOY 情景
END_YEAR = 2070


def _run(tmp_path: Path, monkeypatch) -> tuple[Path, dict[str, Any]]:
    """跑 TOY，返回结果目录与写工作簿时拿到的解和参数。"""
    seen: dict[str, Any] = {}
    real = runner.write_ccs_workbook

    def spy(path, tables, solution, prepared, scenario, assumptions, node_province):
        seen.update(solution=solution, scenario=scenario, assumptions=assumptions)
        real(path, tables, solution, prepared, scenario, assumptions, node_province)

    monkeypatch.setattr(runner, "write_ccs_workbook", spy)
    assert cli.main(["--registry", str(_registry(tmp_path)), *ARGV]) == 0
    return tmp_path / "toy" / "results" / "TOY", seen


def _rows(book, name: str) -> list[list[Any]]:
    return [list(row) for row in book[name].iter_rows(values_only=True)]


def _labelled(rows: list[list[Any]], first: str) -> dict[str, list[Any]]:
    """首列为 *first* 起、到空行为止的一块，按首列取行。"""
    start = next(i for i, row in enumerate(rows) if row[0] == first)
    block: dict[str, list[Any]] = {}
    for row in rows[start + 1:]:
        if row[0] is None:
            break
        block[str(row[0])] = row[1:]
    return block


def test_workbook_reconciles_with_the_tables(tmp_path, monkeypatch) -> None:
    out, seen = _run(tmp_path, monkeypatch)
    book = openpyxl.load_workbook(out / "ccs_results.xlsx")
    tables = {name: pd.read_csv(out / name) for name in TABLES}
    plant, hubs, sinks = tables["plant_detail.csv"], tables["industry_detail.csv"], tables["storage_utilization.csv"]
    assert book.sheetnames == [
        "说明", "Source_Results_2050", "Source_Results_2060", "Sink_Results", "Transmission_Matrix_2050",
        "Transmission_Matrix_2060", *(f"CO2_capture_{k}_stock" for k in ("power", "isi", "cem", "nh3", "meoh",
                                                                         "chemical", "total")),
        "CO2_inject_stock", "num_pipe_stock", "num_pipes_total", "CO2_flow_stock", "GIS_info", "taransport_analysis",
        "Total_CO2_Pipeline_Length", "Cost_Capture", "Cost_Transport", "Cost_Storage", "Revenue_EOR", "Cost_Analysis",
        "objective", "Plant_Pathways", "Industry_Routes",
    ]
    for year in YEARS:
        captured = float(plant.loc[plant["year"] == year, "captured_mt"].sum()
                         + hubs.loc[hubs["year"] == year, "captured_mt"].sum())
        stored = float(sinks.loc[sinks["year"] == year, "storage_use_mtpa"].sum())
        assert captured > 0.5 and stored == pytest.approx(captured, abs=1e-6)  # 前提：toy 在捕集、在封存

        # 各部门捕集之和 = 总捕集，总捕集合计 = 两张明细表的捕集量；投资合计 = 两张明细表的捕集投资。
        rows = _rows(book, f"Source_Results_{year}")
        header = rows[0]
        sources = [dict(zip(header[:18], row[:18])) for row in rows[1:] if row[0] is not None]
        for source in sources:
            sectors = sum(source[f"{s}_Capture"] for s in ("Power", "ISI", "Cement", "Chemical"))
            assert source["Total Capture"] == pytest.approx(sectors)
        assert sum(s["Total Capture"] for s in sources) == pytest.approx(captured)
        capex = tables["plant_cost.csv"].query("year == @year")["ccs_retrofit_capex_cny"].sum() + hubs.query(
            "year == @year")["cost_capital_ccs_cny"].sum()
        assert sum(s["CAPEX_Capture"] for s in sources) * 1e6 == pytest.approx(capex)

        # 输送矩阵：行和 = 该省捕集量、列和 = 该省封存量（toy 全在山西），合计一致。
        matrix = _labelled(_rows(book, f"Transmission_Matrix_{year}"), "源/汇")
        columns = _rows(book, f"Transmission_Matrix_{year}")[0][1:]
        shanxi = dict(zip(columns, matrix["Shanxi"]))
        assert shanxi["Shanxi"] == pytest.approx(captured) and shanxi["合计"] == pytest.approx(captured)
        assert dict(zip(columns, matrix["合计"]))["合计"] == pytest.approx(stored)
        assert all(value is None for key, row in matrix.items() if key not in ("Shanxi", "合计") for value in row)

    # 在役根数 x 档容量 = 在役的管道容量（toy 没有既有走廊容量），工作簿的在役根数与 CSV 相同。
    edges = tables["network_edges.csv"]
    tiers = seen["assumptions"].pipe_capacity_tiers_mtpa
    labels = [f"{t:g}" for t in tiers]
    assert seen["assumptions"].existing_corridor_capacity_mtpa == 0.0
    for row in edges.itertuples(index=False):
        counts = _tier_counts(row.pipes_in_service_by_tier, labels)
        assert sum(c * t for c, t in zip(counts, tiers)) == pytest.approx(row.total_capacity_mtpa, abs=1e-5)
    service = {(r.edge_id, r.year): _tier_counts(r.pipes_in_service_by_tier, labels) for r in edges.itertuples()}
    stock = _rows(book, "num_pipe_stock")[1:]
    assert stock and all(
        (row[4 + k] or 0) == service[(row[0], year)][labels.index(f"{row[3]:g}")]
        for row in stock for k, year in enumerate(YEARS)
    )
    assert (edges["from_province"] == "Shanxi").all() and (edges["to_province"] == "Shanxi").all()

    # 不折现值 x 折现权重 = 进目标函数的值；权重按定义直接写：折现 (1 + r)^-(t - t0)，年度项再乘 10 年一期的年金。
    scenario = seen["scenario"]
    rate, base = scenario.discount_rate, scenario.discount_base_year
    assert rate > 0 and base <= YEARS[0]

    def discount(year: int) -> float:
        return (1.0 + rate) ** -(year - base)

    annuity = (1.0 - (1.0 + rate) ** -10) / rate
    costs = tables["cost_breakdown.csv"]
    for row in costs.itertuples(index=False):
        weight = {"annual": discount(row.year) * annuity, "one_off": discount(row.year),
                  "horizon_end": discount(END_YEAR)}[row.kind]
        assert row.cost_undiscounted_cny * weight == pytest.approx(row.cost_cny, rel=1e-9, abs=1e-6)
    undiscounted = costs.set_index(["category", "year"])["cost_undiscounted_cny"]

    # 残值台账复算出 salvage_credit（最后一年）。
    salvage = ledger_salvage(seen["solution"], YEARS, END_YEAR)
    credit = costs.query("category == 'salvage_credit' and year == 2060")["cost_cny"].iloc[0]
    assert credit < 0 and sum(salvage.values()) == pytest.approx(-credit, rel=1e-9)

    # 成本表：与 cost_breakdown 的同名项相同；封存费 − EOR 抵扣 = storage_cost。
    capture = _labelled(_rows(book, "Cost_Capture"), "百万元")
    transport = _labelled(_rows(book, "Cost_Transport"), "百万元")
    storage = _labelled(_rows(book, "Cost_Storage"), "百万元")["封存费（扣 EOR 抵扣前）"]
    eor = _labelled(_rows(book, "Revenue_EOR"), "百万元")["EOR 抵扣"]
    for k, year in enumerate(YEARS):
        assert capture["煤电捕集岛投资"][k] * 1e6 == pytest.approx(undiscounted[("ccs_retrofit_capex", year)])
        assert capture["煤电捕集岛固定运维"][k] * 1e6 == pytest.approx(undiscounted[("ccs_om_cost", year)])
        assert transport["管道投资"][k] * 1e6 == pytest.approx(undiscounted[("pipe_capex", year)])
        assert transport["运输运维"][k] * 1e6 == pytest.approx(undiscounted[("transport_opex", year)])
        assert (storage[k] - eor[k]) * 1e6 == pytest.approx(undiscounted[("storage_cost", year)])

    # 全期每吨捕集成本按定义复算：捕集投资乘折现、运行费乘折现 x 年金，投资减期末残值，除以捕集量的同权重现值。
    assumptions = seen["assumptions"]
    present = present_captured = 0.0
    for k, year in enumerate(YEARS):
        capex = (capture["煤电捕集岛投资"][k] + capture["工业捕集投资"][k]) * 1e6
        opex = (capture["煤电捕集岛固定运维"][k] + capture["煤电 CCS 额外燃料"][k] + capture["工业捕集年度费"][k]) * 1e6
        present += discount(year) * capex + discount(year) * annuity * opex
        present -= discount(END_YEAR) * remaining_fraction(year, assumptions.ccs_retrofit_lifetime_years, END_YEAR) * (
            capture["煤电捕集岛投资"][k] * 1e6)
        captured = plant.query("year == @year")["captured_mt"].sum() + hubs.query("year == @year")["captured_mt"].sum()
        present_captured += discount(year) * annuity * float(captured) * 1e6
    # toy 的水泥 hub 不建捕集、捕集量为零：上面只扣煤电的残值；捕集量含工业的那部分在部件测试里测（`captured_by_year`）。
    assert capture["工业捕集投资"] == [0, 0]
    horizon = _labelled(_rows(book, "Cost_Analysis"), "全期每吨成本（成本现值 ÷ 捕集量现值）")
    assert horizon["捕集环节"][0] == pytest.approx(present / present_captured, rel=1e-9)

    # 目标函数：各类各年之和 = 目标函数值 = result.json。
    objective = _rows(book, "objective")
    result = json.loads((out.parent / "TOY.json").read_text(encoding="utf-8"))
    assert objective[0][1] * 1e6 == pytest.approx(result["global_objective_cny"])
    assert _labelled(objective, "category")["合计"][-1] == pytest.approx(objective[0][1])

    # 路径份额与明细表相同。
    pathways = _rows(book, "Plant_Pathways")
    ccs = dict(zip(pathways[0], pathways[1]))
    assert [ccs[f"ccs_{y}"] for y in YEARS] == pytest.approx(plant.sort_values("year")["share_ccs"].tolist())


def test_workbook_failure_keeps_the_tables_and_result(tmp_path, monkeypatch, caplog) -> None:
    """工作簿出错只记日志：15 张 CSV 与 result.json 照写，run 正常返回；--force 重跑时上一次的工作簿删掉，不与这次的
    CSV 混放；删不掉（Windows 上在 Excel 里开着）也只记日志。"""
    argv = ["--registry", str(_registry(tmp_path)), *ARGV, "--force"]
    out = tmp_path / "toy" / "results" / "TOY"
    workbook, result = out / "ccs_results.xlsx", out.parent / "TOY.json"
    assert cli.main(argv) == 0 and workbook.exists()

    def broken(*args, **kwargs):
        raise RuntimeError("boom")

    def rerun() -> None:
        """清掉上一次的 CSV 与 result.json 的内容再跑，跑完两样都在，才说明是这一次写的。"""
        for table in out.glob("*.csv"):
            table.unlink()
        result.write_text("{}", encoding="utf-8")
        assert cli.main(argv) == 0
        assert {p.name for p in out.glob("*.csv")} == TABLES
        assert json.loads(result.read_text(encoding="utf-8"))["resolved"]["options"]["force"] is True

    monkeypatch.setattr(runner, "write_ccs_workbook", broken)
    rerun()
    assert not workbook.exists()
    assert "结果工作簿 ccs_results.xlsx 没有写成" in caplog.text and "boom" in caplog.text

    workbook.write_bytes(b"locked")
    real_unlink = Path.unlink

    def locked(path: Path, missing_ok: bool = False) -> None:
        if path.name == workbook.name:
            raise PermissionError(13, "另一个程序正在使用此文件", str(path))
        real_unlink(path, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", locked)
    rerun()
    assert workbook.read_bytes() == b"locked" and "删不掉（[Errno 13] 另一个程序正在使用此文件" in caplog.text
